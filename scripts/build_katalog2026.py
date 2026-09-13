#!/usr/bin/env python3
"""Erzeugt data/katalog2026.json aus data/OPS_Katalogdaten_2026_MASTER.xlsx (Single Source, 1:1-Regel aus CLAUDE.md).

Aufruf:  python3 scripts/build_katalog2026.py [--master PFAD] [--out PFAD] [--check PFAD]
  --master  Excel-Quelle (Standard data/OPS_Katalogdaten_2026_MASTER.xlsx)
  --out     Ziel-JSON (Standard data/katalog2026.json)
  --check   statt zu schreiben: erzeugtes JSON zeichengenau mit dieser Datei vergleichen, Unterschiede als Diff ausgeben

Blätter der Excel:
  Katalog    OPS-Code | Bezeichnung | AOP | Hybrid | Kontextprozedur | Hinweis   ("ja" oder leer). Bitmaske f: 1=AOP, 2=Hybrid, 4=Kontext.
             Hybrid-Vorrang (Regel B. Breuer 02.07.2026): steht AOP und Hybrid auf "ja", zaehlt nur Hybrid.
             Zeilen ohne jedes Flag werden nicht in die JSON uebernommen.
  KX         OPS-Code | Hybrid-DRG   Ausnahmen: Kode ist bei dieser Hybrid-DRG KEINE Kontextprozedur (_KX).
  Kommentar  eine Zeile je _kommentar-Eintrag (Spalte A), Reihenfolge wie im Blatt.
  Meta       Schluessel | Wert   fuer _KATALOG_META (jahr als Zahl, sonst Text).
Ausgabeformat: identisch mit dem Bestand (eine _HD-Zeile je Kode, Reihenfolge = Blattreihenfolge), damit Git-Diffs Zeile fuer Zeile lesbar bleiben.
Keine Werte werden korrigiert, vereinheitlicht oder abgeleitet (oberste Regel CLAUDE.md)."""
import argparse, json, sys, difflib
from pathlib import Path
try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl fehlt: pip install openpyxl")

def ja(v):
    return v is not None and str(v).strip().lower() == "ja"

def lies_master(pfad):
    wb = openpyxl.load_workbook(pfad, data_only=True)
    for blatt in ("Katalog", "KX", "Kommentar", "Meta"):
        if blatt not in wb.sheetnames:
            sys.exit(f"Blatt '{blatt}' fehlt in {pfad}")
    hd = []
    gesehen = set()
    for i, r in enumerate(wb["Katalog"].iter_rows(values_only=True)):
        if i == 0:
            continue
        if not r or r[0] is None:
            continue
        code = str(r[0]).strip()
        if code in gesehen:
            sys.exit(f"Kode doppelt im Blatt Katalog: {code}")
        gesehen.add(code)
        text = "" if r[1] is None else str(r[1]).strip()
        f = (1 if ja(r[2]) else 0) | (2 if ja(r[3]) else 0) | (4 if ja(r[4]) else 0)
        if f & 1 and f & 2:
            f &= ~1  # Hybrid-Vorrang
        if f == 0:
            continue
        hd.append([code, text, f])
    kx = {}
    for i, r in enumerate(wb["KX"].iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None:
            continue
        kx.setdefault(str(r[0]).strip(), []).append(str(r[1]).strip())
    kommentar = [str(r[0]) for r in wb["Kommentar"].iter_rows(values_only=True) if r and r[0] is not None]
    meta = {}
    for i, r in enumerate(wb["Meta"].iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None:
            continue
        k = str(r[0]).strip(); v = r[1]
        meta[k] = int(v) if k == "jahr" else ("" if v is None else str(v))
    return kommentar, meta, kx, hd

def js(o):
    return json.dumps(o, ensure_ascii=False, separators=(",", ":"))

def schreibe(kommentar, meta, kx, hd):
    zeilen = ["{",
              ' "_kommentar": ' + json.dumps(kommentar, ensure_ascii=False) + ",",
              ' "_KATALOG_META": ' + js(meta) + ",",
              ' "_KX": ' + js(kx) + ",",
              ' "_HD": [']
    for i, z in enumerate(hd):
        zeilen.append("  " + js(z) + ("," if i < len(hd) - 1 else ""))
    zeilen += [" ]", "}", ""]
    return "\n".join(zeilen)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--master", default="data/OPS_Katalogdaten_2026_MASTER.xlsx")
    ap.add_argument("--out", default="data/katalog2026.json")
    ap.add_argument("--check", default=None)
    a = ap.parse_args()
    kommentar, meta, kx, hd = lies_master(a.master)
    text = schreibe(kommentar, meta, kx, hd)
    json.loads(text)  # muss gueltig sein
    flags = {}
    for _, _, f in hd:
        flags[f] = flags.get(f, 0) + 1
    print(f"Kodes: {len(hd)} | Flags: {dict(sorted(flags.items()))} | _KX: {len(kx)} | Kommentarzeilen: {len(kommentar)}")
    if a.check:
        alt = Path(a.check).read_text(encoding="utf-8")
        if alt == text:
            print(f"IDENTISCH mit {a.check}")
            return 0
        d = list(difflib.unified_diff(alt.splitlines(), text.splitlines(), a.check, "erzeugt", lineterm="", n=0))
        print(f"UNTERSCHIEDE gegen {a.check}: {sum(1 for l in d if l.startswith('+') and not l.startswith('+++'))} neue, {sum(1 for l in d if l.startswith('-') and not l.startswith('---'))} entfallene Zeilen")
        for l in d[:400]:
            print(l[:160])
        return 1
    Path(a.out).write_text(text, encoding="utf-8")
    print(f"geschrieben: {a.out}")
    return 0

if __name__ == "__main__":
    sys.exit(main())
